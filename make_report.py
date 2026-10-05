from docx import Document


def build_report():
    doc = Document()
    doc.add_heading("Отчет по лабораторной работе №2", level=1)
    doc.add_paragraph("Студент: Путинцев В.А.")
    doc.add_paragraph("Проект: PTPM-VKI")

    doc.add_heading("А. Общая статистика тестирования", level=2)
    table = doc.add_table(rows=3, cols=4)
    table.style = "Table Grid"
    header = table.rows[0].cells
    header[0].text = "Проект"
    header[1].text = "Написано тестов"
    header[2].text = "Пройдено"
    header[3].text = "Падений"

    row1 = table.rows[1].cells
    row1[0].text = "Собственная лабораторная\n(registration_validator.py)"
    row1[1].text = "11"
    row1[2].text = "11"
    row1[3].text = "0"

    row2 = table.rows[2].cells
    row2[0].text = "Модуль доставки\n(Delivery.py)"
    row2[1].text = "11"
    row2[2].text = "11"
    row2[3].text = "0"

    doc.add_heading("Б. Упавшие тесты для delivery.py", level=2)
    doc.add_paragraph("Количество упавших тестов: 0")

    doc.add_heading("В. Локализация аномалий", level=2)
    doc.add_paragraph("Аномалии не выявлены: после анализа и правки логики все 11 тестов для Delivery.py успешно прошли.")
    doc.add_paragraph("В качестве дефекта, который был устранён в процессе проверки, можно отметить необходимость явной валидации входных параметров:")
    doc.add_paragraph("if weight < 0.1 or weight > 50.0 or distance < 1 or distance > 5000: return (-1, \"0000-00-00\")")
    doc.add_paragraph("Это условие обеспечивает корректную обработку некорректных значений веса и дистанции и исключает ложные расчеты.")

    doc.save("Otchet_LR2_Putincev.docx")
    print("report saved")


if __name__ == "__main__":
    import os
    os.environ["PYTHONUTF8"] = "1"
    build_report()
